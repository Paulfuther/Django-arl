import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from arl.user.models import Store

from ..models import (
    SalesImportBatch,
    SalesImportFile,
    SalesImportRow,
    SalesTargetCategory,
    SalesTargetLine,
    SalesTargetPeriod,
)

# print("LOADED SALES IMPORTER:", __file__)

MAX_SALES_FILES = 15


def validate_report_title(worksheet):
    """
    Confirm this is a D365 Category Sales Report.
    """

    title = normalize_text(worksheet["A1"].value).lower()

    if title != "category sales report":
        raise ValueError(
            "This does not appear to be a Category Sales Report. "
            f"Cell A1 contained: {worksheet['A1'].value!r}"
        )


def normalize_text(value):
    if value is None:
        return ""

    return str(value).strip()


def normalize_code(value):
    """
    Normalize a D365 category code.

    Examples:
        2310 -> "2310"
        "2310" -> "2310"
        "2310.0" -> "2310"
    """

    if value is None:
        return ""

    text = str(value).strip()

    if text.endswith(".0"):
        text = text[:-2]

    return text


def parse_category_code(value):
    """
    Convert:

        2310 - ALCOHOLIC BEVERAGES

    into:

        2310
    """

    text = normalize_text(value)

    match = re.match(
        r"^(\d+)\s*-\s*",
        text,
    )

    if not match:
        return None

    return match.group(1)


def parse_decimal(value):
    """
    Convert spreadsheet sales values into Decimal.
    """

    if value in (None, ""):
        return Decimal("0.00")

    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"))

    if isinstance(value, (int, float)):
        return Decimal(str(value)).quantize(Decimal("0.01"))

    cleaned_value = (
        str(value)
        .replace("$", "")
        .replace(",", "")
        .replace("(", "-")
        .replace(")", "")
        .strip()
    )

    try:
        return Decimal(cleaned_value).quantize(Decimal("0.01"))

    except InvalidOperation as exc:
        raise ValueError(f"Invalid sales amount: {value}") from exc


def parse_excel_date(value):
    """
    Read an Excel date or text date.
    """

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):
        text = value.strip()

        date_formats = (
            "%Y-%m-%d",
            "%m/%d/%Y",
            "%Y/%m/%d",
            "%d/%m/%Y",
        )

        for date_format in date_formats:
            try:
                return datetime.strptime(
                    text,
                    date_format,
                ).date()

            except ValueError:
                continue

    raise ValueError(f"Unable to read date: {value}")


def get_import_categories(employer):
    """
    Return all active target categories that have
    a D365 category code.
    """

    return list(
        SalesTargetCategory.objects.filter(
            employer=employer,
            is_active=True,
        )
        .exclude(
            d365_category_code__isnull=True,
        )
        .exclude(
            d365_category_code="",
        )
        .order_by("name")
    )


def find_store_number(worksheet):
    """
    Find the store number.

    Expected example:
        65077 - 1553690 ONTARIO INC.

    Works whether the heading is merged across A:I
    or stored only in cell A3.
    """

    value = worksheet["A3"].value
    text = normalize_text(value)

    match = re.match(
        r"^\s*(\d+)\s*-\s*",
        text,
    )

    if match:
        return match.group(1)

    # Fallback: search first 15 rows and columns A:I.
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=15,
        min_col=1,
        max_col=9,
    ):
        for cell in row:
            text = normalize_text(cell.value)

            match = re.match(
                r"^\s*(\d+)\s*-\s*",
                text,
            )

            if match:
                return match.group(1)

    raise ValueError(
        "Store number could not be found in cell A3 or within the first 15 rows."
    )


def find_report_date(
    worksheet,
    label,
):
    """
    Find report dates such as:

        From date    2026-07-01
        To date      2026-07-31

    Searches columns A:I in the first 20 rows.
    """

    target_label = label.strip().lower()

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=20,
        min_col=1,
        max_col=9,
    ):
        row_cells = list(row)

        for index, cell in enumerate(row_cells):
            cell_text = normalize_text(cell.value).lower()

            if cell_text != target_label:
                continue

            # Find the next nonblank value on the same row.
            for next_cell in row_cells[index + 1 :]:
                if next_cell.value not in (None, ""):
                    return parse_excel_date(next_cell.value)

    raise ValueError(f"{label} could not be found.")


def find_header_row(worksheet):
    """
    Find the row containing 'Category name'.

    Only columns A:I need to be examined.
    """

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            min_col=1,
            max_col=9,
        ),
        start=1,
    ):
        for cell in row:
            if normalize_text(cell.value).lower() == "category name":
                return row_number

    raise ValueError("The Category name heading could not be found.")


def find_column_number(
    worksheet,
    header_row,
    heading_name,
):
    """
    Find a heading within columns A:I.
    """

    target_heading = heading_name.strip().lower()

    row = next(
        worksheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            min_col=1,
            max_col=9,
        ),
        None,
    )

    if row is None:
        raise ValueError(f"Header row {header_row} could not be read.")

    for column_number, cell in enumerate(
        row,
        start=1,
    ):
        heading = normalize_text(cell.value).lower()

        if heading == target_heading:
            return column_number

    raise ValueError(
        f"The {heading_name} column could not be found in columns A through I."
    )


def read_category_sales(
    worksheet,
    expected_codes,
):
    """
    Read only the configured D365 category rows.

    The report uses columns A:I:
        A = Category name
        D = Sales amount

    The column headings are still detected dynamically.
    """

    header_row = find_header_row(worksheet)

    category_column = find_column_number(
        worksheet=worksheet,
        header_row=header_row,
        heading_name="Category name",
    )

    sales_column = find_column_number(
        worksheet=worksheet,
        header_row=header_row,
        heading_name="Sales amount",
    )

    expected_codes = {normalize_code(code) for code in expected_codes}

    category_sales = {}
    duplicate_codes = []

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        min_col=1,
        max_col=9,
    ):
        category_value = row[category_column - 1].value

        category_code = parse_category_code(category_value)

        if category_code not in expected_codes:
            continue

        if category_code in category_sales:
            duplicate_codes.append(category_code)
            continue

        sales_value = row[sales_column - 1].value

        category_sales[category_code] = parse_decimal(sales_value)

    if duplicate_codes:
        duplicate_text = ", ".join(sorted(set(duplicate_codes)))

        raise ValueError(
            f"Duplicate configured category codes were found: {duplicate_text}"
        )

    return category_sales


def get_store(employer, store_number):
    return Store.objects.get(
        employer=employer,
        number=store_number,
    )


def get_period(
    employer,
    from_date,
    to_date,
):
    return SalesTargetPeriod.objects.get(
        employer=employer,
        start_date=from_date,
        end_date=to_date,
    )


def parse_and_stage_file(
    *,
    batch,
    uploaded_file,
):
    """
    Parse one spreadsheet and create staging rows.

    This does not update SalesTargetLine.
    It only prepares the preview.
    """

    # print("PARSE FUNCTION CALLED:", uploaded_file.name)

    staged_file = SalesImportFile.objects.create(
        batch=batch,
        original_filename=uploaded_file.name,
    )

    workbook = None

    try:
        # -----------------------------------------
        # Load configured import categories
        # -----------------------------------------

        import_categories = get_import_categories(batch.employer)

        if not import_categories:
            raise ValueError(
                "No active sales target categories have "
                "a D365 category code configured."
            )

        categories_by_code = {
            normalize_code(category.d365_category_code): category
            for category in import_categories
        }

        # print(
        #    "CATEGORY CODES:",
        #    list(categories_by_code.keys()),
        # )

        # -----------------------------------------
        # Open workbook
        # -----------------------------------------

        workbook = load_workbook(
            uploaded_file,
            data_only=True,
        )

        worksheet = workbook.active
        validate_report_title(worksheet)

        # -----------------------------------------
        # Read and save store number
        # -----------------------------------------

        store_number = find_store_number(worksheet)

        # print(
        #    "STORE NUMBER RETURNED:",
        #    store_number,
        # )

        staged_file.store_number = store_number

        staged_file.save(
            update_fields=[
                "store_number",
            ]
        )

        # -----------------------------------------
        # Read and save reporting dates
        # -----------------------------------------

        from_date = find_report_date(
            worksheet,
            "From date",
        )

        to_date = find_report_date(
            worksheet,
            "To date",
        )

        staged_file.from_date = from_date
        staged_file.to_date = to_date

        staged_file.save(
            update_fields=[
                "from_date",
                "to_date",
            ]
        )

        # -----------------------------------------
        # Read configured category sales
        # -----------------------------------------

        category_sales = read_category_sales(
            worksheet=worksheet,
            expected_codes=categories_by_code.keys(),
        )

        # -----------------------------------------
        # Validate store
        # -----------------------------------------

        try:
            store = get_store(
                employer=batch.employer,
                store_number=store_number,
            )

        except Store.DoesNotExist as exc:
            raise ValueError(
                f"Store {store_number} was not found for this employer."
            ) from exc

        except Store.MultipleObjectsReturned as exc:
            raise ValueError(
                f"Multiple store records were found for store number {store_number}."
            ) from exc

        # -----------------------------------------
        # Validate selected target period
        # -----------------------------------------

        period = batch.target_period

        if period.employer_id != batch.employer_id:
            raise ValueError(
                "The selected target period does not belong to this employer."
            )

        if period.start_date != from_date or period.end_date != to_date:
            raise ValueError(
                "The spreadsheet dates do not match the selected "
                f"sales target period '{period.name}'. "
                f"Spreadsheet: {from_date} to {to_date}. "
                f"Selected period: {period.start_date} "
                f"to {period.end_date}."
            )

        # -----------------------------------------
        # Check for missing configured categories
        # -----------------------------------------

        missing_codes = [
            code for code in categories_by_code if code not in category_sales
        ]

        if missing_codes:
            missing_categories = [
                (f"{code} - {categories_by_code[code].name}") for code in missing_codes
            ]

            raise ValueError(
                "The following configured categories were "
                "not found in the spreadsheet: " + ", ".join(missing_categories)
            )

        # -----------------------------------------
        # Create preview rows
        # -----------------------------------------

        file_has_errors = False

        for category_code, target_category in categories_by_code.items():
            imported_sales = category_sales[category_code]

            try:
                target_line = SalesTargetLine.objects.get(
                    period=period,
                    store=store,
                    category=target_category,
                )

            except SalesTargetLine.DoesNotExist:
                # This store/campaign does not track this category.
                # Ignore it completely.
                continue

            except SalesTargetLine.MultipleObjectsReturned:
                file_has_errors = True

                SalesImportRow.objects.create(
                    staged_file=staged_file,
                    target_category=target_category,
                    category_code=category_code,
                    category_name=target_category.name,
                    imported_sales=imported_sales,
                    current_sales=Decimal("0.00"),
                    difference=imported_sales,
                    is_valid=False,
                    error_message=(
                        "Multiple sales target lines exist for "
                        f"store {store_number} and "
                        f"{target_category.name}."
                    ),
                )

                continue

            current_sales = target_line.current_sales or Decimal("0.00")

            SalesImportRow.objects.create(
                staged_file=staged_file,
                target_category=target_category,
                category_code=category_code,
                category_name=target_category.name,
                imported_sales=imported_sales,
                current_sales=current_sales,
                difference=(imported_sales - current_sales),
                is_valid=True,
                error_message="",
            )

        if not staged_file.rows.exists() and not file_has_errors:
            raise ValueError(
                f"No sales targets were found for store "
                f"{store_number} in campaign '{period.name}'."
            )
        # -----------------------------------------
        # Mark file ready or invalid
        # -----------------------------------------

        staged_file.is_valid = not file_has_errors

        if file_has_errors:
            staged_file.error_message = (
                "One or more duplicate sales target lines were found."
            )
        else:
            staged_file.error_message = ""

        staged_file.save(
            update_fields=[
                "is_valid",
                "error_message",
            ]
        )

    except Exception as exc:
        # traceback.print_exc()

        error_text = str(exc)

        # print(
        #    "SALES IMPORT ERROR:",
        #    error_text,
        # )

        staged_file.is_valid = False
        staged_file.error_message = error_text

        staged_file.save(
            update_fields=[
                "is_valid",
                "error_message",
            ]
        )

    finally:
        if workbook is not None:
            workbook.close()

    return staged_file


def create_preview_batch(
    *,
    employer,
    uploaded_by,
    uploaded_files,
    target_period,
):
    if not uploaded_files:
        raise ValueError("Please select at least one spreadsheet.")

    if len(uploaded_files) > MAX_SALES_FILES:
        raise ValueError(f"You can upload a maximum of {MAX_SALES_FILES} spreadsheets.")

    if target_period.employer_id != employer.id:
        raise ValueError("The selected sales period does not belong to your employer.")

    invalid_filenames = [
        uploaded_file.name
        for uploaded_file in uploaded_files
        if not uploaded_file.name.lower().endswith((".xlsx", ".xlsm"))
    ]

    if invalid_filenames:
        raise ValueError(
            "Only .xlsx and .xlsm files are allowed. "
            "Invalid files: " + ", ".join(invalid_filenames)
        )

    batch = SalesImportBatch.objects.create(
        employer=employer,
        uploaded_by=uploaded_by,
        target_period=target_period,
    )

    for uploaded_file in uploaded_files:
        parse_and_stage_file(
            batch=batch,
            uploaded_file=uploaded_file,
        )

    return batch


def validate_batch(batch):
    """
    Validate every spreadsheet and return useful error details.
    """

    files = list(batch.files.prefetch_related("rows").all())

    if not files:
        return False, "No spreadsheets were uploaded."

    if batch.status != SalesImportBatch.STATUS_PREVIEW:
        return (
            False,
            "This import batch is no longer available.",
        )

    file_errors = []

    for staged_file in files:
        if staged_file.is_valid:
            continue

        filename = staged_file.original_filename
        store_number = staged_file.store_number or "Unknown store"

        error_message = staged_file.error_message or "Unknown spreadsheet error."

        file_errors.append(f"{filename} ({store_number}): {error_message}")

    if file_errors:
        return (
            False,
            " | ".join(file_errors),
        )

    invalid_rows = SalesImportRow.objects.filter(
        staged_file__batch=batch,
        is_valid=False,
    ).select_related(
        "staged_file",
        "target_category",
    )

    row_errors = []

    for staged_row in invalid_rows:
        row_errors.append(
            (
                f"Store "
                f"{staged_row.staged_file.store_number}: "
                f"{staged_row.category_name} - "
                f"{staged_row.error_message}"
            )
        )

    if row_errors:
        return (
            False,
            " | ".join(row_errors),
        )

    store_numbers = [staged_file.store_number for staged_file in files]

    duplicate_store_numbers = sorted(
        {
            store_number
            for store_number in store_numbers
            if store_numbers.count(store_number) > 1
        }
    )

    if duplicate_store_numbers:
        return (
            False,
            "Duplicate spreadsheets were uploaded for "
            "these stores: " + ", ".join(duplicate_store_numbers),
        )

    report_periods = {
        (
            staged_file.from_date,
            staged_file.to_date,
        )
        for staged_file in files
    }

    if len(report_periods) != 1:
        return (
            False,
            "All spreadsheets must use the same From date and To date.",
        )

    row_counts = {staged_file.rows.count() for staged_file in files}

    if len(row_counts) != 1:
        return (
            False,
            "The spreadsheets do not contain the same "
            "number of configured sales categories.",
        )

    return True, ""


@transaction.atomic
def commit_sales_import(batch):
    """
    Update all staged sales values in one transaction.

    If any update fails, all changes are rolled back.
    """

    locked_batch = SalesImportBatch.objects.select_for_update().get(pk=batch.pk)

    is_valid, validation_error = validate_batch(locked_batch)

    if not is_valid:
        raise ValueError(validation_error)

    staged_files = locked_batch.files.prefetch_related("rows__target_category").all()

    updated_count = 0

    # Loop over every store spreadsheet.
    for staged_file in staged_files:
        store = get_store(
            employer=locked_batch.employer,
            store_number=staged_file.store_number,
        )

        period = locked_batch.target_period

        # Loop over every configured category.
        for staged_row in staged_file.rows.all():
            if staged_row.is_skipped:
                continue

            target_line = SalesTargetLine.objects.select_for_update().get(
                period=period,
                store=store,
                category=staged_row.target_category,
            )

            target_line.current_sales = staged_row.imported_sales

            target_line.save(update_fields=["current_sales"])

            updated_count += 1

    locked_batch.status = SalesImportBatch.STATUS_IMPORTED

    locked_batch.imported_at = timezone.now()

    locked_batch.save(
        update_fields=[
            "status",
            "imported_at",
        ]
    )

    return updated_count
