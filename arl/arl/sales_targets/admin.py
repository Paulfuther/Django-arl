from django.contrib import admin, messages
from .models import SalesTargetCategory, SalesTargetPeriod, SalesTargetLine
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import path
from django.utils.html import format_html
from django.urls import reverse

def export_sales_target_management_report(modeladmin, request, queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Management Report"

    for period in queryset:

        # Report title
        ws.append(["Sales Target Management Report"])
        ws["A1"].font = Font(size=16, bold=True)

        # Period information
        ws.append([f"Period: {period.name}"])
        ws["A2"].font = Font(bold=True)

        ws.append([f"Dates: {period.start_date} to {period.end_date}"])
        ws["A3"].font = Font(bold=True)

        ws.append([])

        # Column headings
        headers = [
            "Store",
            "Category",
            "Target Amount",
            "Current Sales",
            "% To Target",
            "Projected Sales",
            "Projected Variance",
            "Required Daily Sales",
            "Status",
        ]

        ws.append(headers)

        header_row = ws.max_row

        for cell in ws[header_row]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        previous_store = None

        lines = (
            period.target_lines
            .select_related("store", "category", "period")
            .order_by("store_id", "category__name")
        )

        for line in lines:

            # Insert a blank row when the store changes
            if previous_store is not None and previous_store != line.store_id:
                ws.append([])

            ws.append([
                line.store.number,
                line.category.name,
                round(float(line.target_amount)),
                round(float(line.current_sales)),
                round(float(line.percent_to_target)),
                round(float(line.projected_sales)),
                round(float(line.projected_variance)),
                round(float(line.required_daily_sales)),
                line.status,
            ])

            row = ws.max_row

            # Projected Variance (Column G = 7)
            variance_cell = ws.cell(row=row, column=7)

            if line.projected_variance > 0:
                variance_cell.font = Font(color="008000", bold=True)
            elif line.projected_variance < 0:
                variance_cell.font = Font(color="FF0000", bold=True)

            # Required Daily Sales (Column H = 8)
            ws.cell(row=row, column=8).number_format = "#,##0"

            # Status (Column I = 9)
            status_cell = ws.cell(row=row, column=9)

            # Keep status on one line
            status_cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=False,
            )

            # Replace emoji circles with consistent-width circles
            if "Ahead" in line.status:
                status_cell.value = "● Ahead"
                status_cell.font = Font(color="008000", bold=True)

            elif "On Track" in line.status:
                status_cell.value = "● On Track"
                status_cell.font = Font(color="C9A000", bold=True)

            elif "Behind" in line.status:
                status_cell.value = "● Behind"
                status_cell.font = Font(color="C00000", bold=True)

            else:
                status_cell.value = "○ Not Started"
                status_cell.font = Font(color="808080", bold=True)

            previous_store = line.store_id

    # Auto-size columns
    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter

        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(max_length + 2, 45)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        'attachment; filename="sales_target_management_report.xlsx"'
    )

    wb.save(response)
    return response


class SalesTargetLineInline(admin.TabularInline):
    model = SalesTargetLine
    extra = 1
    fields = (
        "store",
        "category",
        "target_amount",
        "current_sales",
    )
    
    readonly_fields = ()
    can_delete = False


@admin.register(SalesTargetPeriod)
class SalesTargetPeriodAdmin(admin.ModelAdmin):
    list_display = (
        "name",
        "start_date",
        "end_date",
        "update_sales_link",
    )

    inlines = [SalesTargetLineInline]
    actions = [
        export_sales_target_management_report,
    ]
    def update_sales_link(self, obj):
        url = reverse(
            "admin:sales_target_update_current_sales",
            args=[obj.id],
        )
        return format_html(
            '<a class="button" href="{}">Update Current Sales</a>',
            url,
        )

    update_sales_link.short_description = "Current Sales"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                "<int:period_id>/update-current-sales/",
                self.admin_site.admin_view(self.update_current_sales_view),
                name="sales_target_update_current_sales",
            ),
        ]
        return custom_urls + urls

    def update_current_sales_view(self, request, period_id):
        period = SalesTargetPeriod.objects.get(id=period_id)

        lines = SalesTargetLine.objects.filter(
            period=period
        ).select_related(
            "store",
            "category",
        ).order_by(
            "store__number",
            "category__name",
        )

        if request.method == "POST":
            for line in lines:
                value = request.POST.get(f"current_sales_{line.id}", "0").strip()

                if value == "":
                    value = "0"

                line.current_sales = value
                line.save(update_fields=["current_sales"])

            messages.success(request, "Current sales updated successfully.")
            return redirect("..")

        context = {
            **self.admin_site.each_context(request),
            "period": period,
            "lines": lines,
            "title": f"Update Current Sales - {period.name}",
        }

        return render(
            request,
            "admin/sales_targets/update_current_sales.html",
            context,
        )


@admin.register(SalesTargetCategory)
class SalesTargetCategoryAdmin(admin.ModelAdmin):
    list_display = ("name", "employer", "is_active")
    list_filter = ("employer", "is_active")
    search_fields = ("name", "employer__name")


@admin.register(SalesTargetLine)
class SalesTargetLineAdmin(admin.ModelAdmin):
    list_display = (
        "period",
        "store_number",
        "category",
        "target_amount",
        "current_sales",
        "formatted_percent_to_target",
        "status",
        "formatted_projected_sales",
        "formatted_projected_variance",
        "required_daily_sales_display",
    )
    def store_number(self, obj):

        return obj.store.number

    store_number.short_description = "Store"

    store_number.admin_order_field = "store__number"

    def required_daily_sales_display(self, obj):

        return int(obj.required_daily_sales)

    required_daily_sales_display.short_description = "Required Daily Sales"

    list_filter = ("period", "store", "category")
    search_fields = (
        "period__name",
        "store__name",
        "category__name",
    )
    fields = (
        "store",
        "category",
        "target_amount",
        "current_sales",
        "last_updated",
    )

    def formatted_percent_to_target(self, obj):
        return f"{obj.percent_to_target:.1f}%"
    formatted_percent_to_target.short_description = "% to Target"


    def formatted_projected_sales(self, obj):
        return f"{obj.projected_sales:,.2f}"
    formatted_projected_sales.short_description = "Projected Sales"


    def formatted_projected_variance(self, obj):
        return f"{obj.projected_variance:,.2f}"
    formatted_projected_variance.short_description = "Projected Variance"
