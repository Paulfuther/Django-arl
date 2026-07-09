import json
from collections import defaultdict
from decimal import Decimal

from django.conf import settings
from django.db.models import Max
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import SalesTargetLine, SalesTargetPeriod


def index(request):
    return HttpResponse("Sales targets app is working.")


def build_morning_summary(period):
    lines = period.target_lines.select_related("store", "category")

    # Top 10 worst issues overall
    behind_lines = sorted(
        [line for line in lines if line.projected_variance < 0],
        key=lambda line: line.projected_variance,
    )[:10]

    # Display those top 10 grouped by store
    behind_lines = sorted(behind_lines, key=lambda line: line.store.number)

    message = f"Morning Sales Target Update - {period.name}\n\n"
    message += "Top 10 Opportunities:\n"

    current_store = None

    for line in behind_lines:
        if current_store != line.store.number:
            current_store = line.store.number
            message += f"\nStore {current_store}\n"

        message += (
            f"  - {line.category.name}: "
            f"${round(abs(line.projected_variance)):,} behind. "
            f"Needs ${round(line.required_daily_sales):,}/day.\n"
        )

    return message


def build_today_focus(period):
    lines = period.target_lines.select_related("store", "category")

    top_issues = sorted(
        [line for line in lines if line.projected_variance < 0],
        key=lambda line: line.projected_variance,
    )[:10]

    by_store = defaultdict(list)

    for line in top_issues:
        by_store[line.store.number].append(line.category.name)

    focus = []

    for store_number, categories in sorted(by_store.items()):
        if len(categories) == 1:
            focus.append(f"Store {store_number} needs attention in {categories[0]}.")
        else:
            focus.append(
                f"Store {store_number} needs attention in {', '.join(categories[:-1])} and {categories[-1]}."
            )

    return focus


def build_biggest_win(period):
    lines = period.target_lines.select_related("store", "category")

    ahead_lines = [line for line in lines if line.projected_variance > 0]

    if not ahead_lines:
        return None

    return max(ahead_lines, key=lambda line: line.projected_variance)


def sales_target_dashboard(request, period_id):
    period = get_object_or_404(SalesTargetPeriod, id=period_id)

    lines = (
        SalesTargetLine.objects.filter(period=period)
        .select_related("store", "category")
        .order_by("store__number", "category__name")
    )

    total_target = Decimal("0")
    total_current = Decimal("0")
    total_projected = Decimal("0")

    needs_attention = []

    category_totals = defaultdict(
        lambda: {
            "target": Decimal("0"),
            "current": Decimal("0"),
            "projected": Decimal("0"),
        }
    )

    for line in lines:
        total_target += line.target_amount
        total_current += line.current_sales
        total_projected += line.projected_sales

        category_name = line.category.name
        category_totals[category_name]["target"] += line.target_amount
        category_totals[category_name]["current"] += line.current_sales
        category_totals[category_name]["projected"] += line.projected_sales

        if line.projected_variance < 0:
            needs_attention.append(
                {
                    "store": line.store.number,
                    "category": line.category.name,
                    "variance": line.projected_variance,
                    "required_daily_sales": line.required_daily_sales,
                    "status": line.status,
                }
            )

    projected_variance = total_projected - total_target

    category_summary = []

    for category, totals in category_totals.items():
        variance = totals["projected"] - totals["target"]

        if variance > 0:
            status = "🟢 Ahead"
        elif variance >= -(totals["target"] * Decimal("0.05")):
            status = "🟡 On Track"
        else:
            status = "🔴 Behind"

        category_summary.append(
            {
                "category": category,
                "target": totals["target"],
                "current": totals["current"],
                "projected": totals["projected"],
                "variance": variance,
                "status": status,
            }
        )

    needs_attention = sorted(needs_attention, key=lambda x: x["variance"])[:10]
    category_summary = sorted(category_summary, key=lambda x: x["variance"])
    morning_summary = build_morning_summary(period)
    today_focus = build_today_focus(period)
    biggest_win = build_biggest_win(period)
    last_updated = period.target_lines.aggregate(
        Max("updated_at")
    )["updated_at__max"]
    target_lines = period.target_lines.select_related(
        "store",
        "category"
    ).order_by(
        "store__number",
        "category__name"
    )

    context = {
        "period": period,
        "total_target": total_target,
        "total_current": total_current,
        "total_projected": total_projected,
        "projected_variance": projected_variance,
        "needs_attention": needs_attention,
        "category_summary": category_summary,
        "morning_summary": morning_summary,
        "today_focus": today_focus,
        "biggest_win": biggest_win,
        "last_updated": last_updated,
        "target_lines": target_lines,
    }

    return render(
        request,
        "sales_targets/dashboard.html",
        context,
    )


def sales_target_campaign_selector(request):
    periods = SalesTargetPeriod.objects.all().order_by("-start_date")

    selected_period_id = request.GET.get("period")
    selected_period = None
    target_lines = None

    if selected_period_id:
        selected_period = get_object_or_404(SalesTargetPeriod, id=selected_period_id)

        target_lines = (
            SalesTargetLine.objects.filter(period=selected_period)
            .select_related("store", "category")
            .order_by("store__store_number", "category__name")
        )

    return render(
        request,
        "sales_targets/campaign_selector.html",
        {
            "periods": periods,
            "selected_period": selected_period,
            "target_lines": target_lines,
        },
    )


def sales_target_dashboard_selector(request):
    periods = SalesTargetPeriod.objects.all().order_by("-start_date")

    if request.method == "POST":
        period_id = request.POST.get("period_id")

        if period_id:
            return redirect("sales_targets:sales_target_dashboard", period_id=period_id)

    return render(
        request,
        "sales_targets/dashboard_selector.html",
        {
            "periods": periods,
        },
    )


def export_sales_target_summary_dashboard(request, period_id):
    period = get_object_or_404(SalesTargetPeriod, id=period_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Summary"

    summary_fill = PatternFill(
        start_color="D9EAF7",
        end_color="D9EAF7",
        fill_type="solid",
    )

    lines = period.target_lines.select_related("store", "category", "period").order_by(
        "store_id", "category__name"
    )

    total_target = sum(line.target_amount for line in lines)
    total_current = sum(line.current_sales for line in lines)
    total_projected = sum(line.projected_sales for line in lines)
    projected_variance = total_projected - total_target

    ws.append(["Sales Target Summary Report"])
    ws["A1"].font = Font(size=16, bold=True)

    ws.append([f"Period: {period.name}"])
    ws.append([f"Dates: {period.start_date} to {period.end_date}"])
    ws.append([])

    ws.append(["Executive Summary"])
    ws[ws.max_row][0].font = Font(bold=True)

    ws.append(["Total Target", round(float(total_target))])
    ws.append(["Current Sales", round(float(total_current))])
    ws.append(["Projected Sales", round(float(total_projected))])
    ws.append(["Projected Variance", round(float(projected_variance))])
    ws.append([])

    for row in range(ws.max_row - 4, ws.max_row):
        for col in (1, 2):
            cell = ws.cell(row=row, column=col)
            cell.fill = summary_fill
            cell.font = Font(bold=True)

    ws.append(["Current Opportunities"])
    ws[ws.max_row][0].font = Font(bold=True)

    ws.append(
        [
            "Store",
            "Category",
            "Projected Variance",
            "Required Daily Sales",
            "Status",
        ]
    )

    attention_lines = sorted(
        [line for line in lines if line.projected_variance < 0],
        key=lambda line: line.projected_variance,
    )[:10]

    for line in attention_lines:
        ws.append(
            [
                line.store.number,
                line.category.name,
                round(float(line.projected_variance)),
                round(float(line.required_daily_sales)),
                line.status,
            ]
        )

    ws.append([])

    ws.append(["Category Summary"])
    ws[ws.max_row][0].font = Font(bold=True)

    ws.append(
        [
            "Category",
            "Target",
            "Current Sales",
            "Projected Sales",
            "Projected Variance",
            "Status",
        ]
    )

    category_totals = {}

    for line in lines:
        category = line.category.name

        if category not in category_totals:
            category_totals[category] = {
                "target": Decimal("0"),
                "current": Decimal("0"),
                "projected": Decimal("0"),
            }

        category_totals[category]["target"] += line.target_amount
        category_totals[category]["current"] += line.current_sales
        category_totals[category]["projected"] += line.projected_sales

    for category, totals in category_totals.items():
        variance = totals["projected"] - totals["target"]

        if variance > 0:
            status = "● Ahead"
        elif variance >= -(totals["target"] * Decimal("0.05")):
            status = "● On Track"
        else:
            status = "● Behind"

        ws.append(
            [
                category,
                round(float(totals["target"])),
                round(float(totals["current"])),
                round(float(totals["projected"])),
                round(float(variance)),
                status,
            ]
        )

        row = ws.max_row
        variance_cell = ws.cell(row=row, column=5)
        status_cell = ws.cell(row=row, column=6)

        if variance > 0:
            variance_cell.font = Font(color="008000", bold=True)
            status_cell.font = Font(color="008000", bold=True)
        elif variance < 0:
            variance_cell.font = Font(color="C00000", bold=True)
            status_cell.font = Font(color="C00000", bold=True)

    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter

        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(max_length + 2, 35)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="sales_target_summary_{period.name}.xlsx"'
    )

    wb.save(response)
    return response


def export_sales_target_management_dashboard(request, period_id):
    period = get_object_or_404(SalesTargetPeriod, id=period_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Management Report"

    # Report title
    ws.append(["Sales Target Management Report"])
    ws["A1"].font = Font(size=16, bold=True)

    # Period information
    ws.append([f"Period: {period.name}"])
    ws["A2"].font = Font(bold=True)

    ws.append([f"Dates: {period.start_date} to {period.end_date}"])
    ws["A3"].font = Font(bold=True)

    ws.append([])

    headers = [
        "Store",
        "Category",
        "Target Amount",
        "Current Sales",
        "% To Target",
        "Projected Sales",
        "Projected Variance",
        "Required Daily Sales",
        "Status",
    ]

    ws.append(headers)

    header_row = ws.max_row

    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    previous_store = None

    lines = period.target_lines.select_related("store", "category", "period").order_by(
        "store_id", "category__name"
    )

    for line in lines:
        if previous_store is not None and previous_store != line.store_id:
            ws.append([])

        ws.append(
            [
                line.store.number,
                line.category.name,
                round(float(line.target_amount)),
                round(float(line.current_sales)),
                round(float(line.percent_to_target)),
                round(float(line.projected_sales)),
                round(float(line.projected_variance)),
                round(float(line.required_daily_sales)),
                line.status,
            ]
        )

        row = ws.max_row

        variance_cell = ws.cell(row=row, column=7)

        if line.projected_variance > 0:
            variance_cell.font = Font(color="008000", bold=True)
        elif line.projected_variance < 0:
            variance_cell.font = Font(color="FF0000", bold=True)

        ws.cell(row=row, column=8).number_format = "#,##0"

        status_cell = ws.cell(row=row, column=9)

        status_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=False,
        )

        if "Ahead" in line.status:
            status_cell.value = "● Ahead"
            status_cell.font = Font(color="008000", bold=True)

        elif "On Track" in line.status:
            status_cell.value = "● On Track"
            status_cell.font = Font(color="C9A000", bold=True)

        elif "Behind" in line.status:
            status_cell.value = "● Behind"
            status_cell.font = Font(color="C00000", bold=True)

        else:
            status_cell.value = "○ Not Started"
            status_cell.font = Font(color="808080", bold=True)

        previous_store = line.store_id

    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter

        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(max_length + 2, 45)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="sales_target_management_{period.name}.xlsx"'
    )

    wb.save(response)
    return response


def generate_ai_sales_coaching(period):
    client = OpenAI(api_key=settings.OPENAI_API_KEY)

    lines = period.target_lines.select_related("store", "category")

    sales_data = []

    for line in lines:
        sales_data.append(
            {
                "store": line.store.number,
                "category": line.category.name,
                "target": round(float(line.target_amount)),
                "current_sales": round(float(line.current_sales)),
                "projected_sales": round(float(line.projected_sales)),
                "projected_variance": round(float(line.projected_variance)),
                "required_daily_sales": round(float(line.required_daily_sales)),
                "status": line.status,
            }
        )

    prompt = f"""
        You are a retail sales performance coach.

        Analyze this sales target campaign.

        Campaign:
        {period.name}
        Dates:
        {period.start_date} to {period.end_date}

        Data:
        {json.dumps(sales_data, indent=2)}

        Write:
        1. A short executive summary.
        2. The top 3 priorities for today.
        3. Practical coaching suggestions for managers.
        4. One WhatsApp-ready message.

        Rules:
        - Do not invent numbers.
        - Only use the data provided.
        - Be direct, practical, and concise.
        - Keep the WhatsApp message friendly and professional.
        """

    response = client.responses.create(
        model="gpt-4.1-mini",
        input=prompt,
    )

    return response.output_text


def ai_sales_coaching_dashboard(request, period_id):
    period = get_object_or_404(SalesTargetPeriod, id=period_id)

    ai_coaching = generate_ai_sales_coaching(period)

    return render(
        request,
        "sales_targets/ai_sales_coaching.html",
        {
            "period": period,
            "ai_coaching": ai_coaching,
        },
    )
